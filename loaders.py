from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document as DocxDocument
from langchain_core.documents import Document
from pypdf import PdfReader


TEXT_EXTENSIONS = {
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".java",
    ".kt",
    ".go",
    ".rs",
    ".cpp",
    ".c",
    ".cs",
    ".sql",
    ".md",
    ".txt",
    ".json",
    ".yaml",
    ".yml",
    ".toml",
    ".ini",
    ".sh",
    ".bat",
    ".ps1",
    ".html",
    ".css",
}


def _read_text(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp949", "euc-kr", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(errors="ignore")


def _load_pdf(path: Path) -> list[Document]:
    reader = PdfReader(str(path))
    docs: list[Document] = []
    for idx, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if not text.strip():
            continue
        docs.append(
            Document(
                page_content=text,
                metadata={"source": path.name, "page": idx, "type": "pdf"},
            )
        )
    return docs


def _load_docx(path: Path) -> list[Document]:
    doc = DocxDocument(str(path))
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip())
    if not text.strip():
        return []
    return [
        Document(
            page_content=text,
            metadata={"source": path.name, "type": "docx"},
        )
    ]


def _load_text(path: Path) -> list[Document]:
    text = _read_text(path)
    if not text.strip():
        return []
    return [
        Document(
            page_content=text,
            metadata={"source": path.name, "type": "text", "extension": path.suffix},
        )
    ]


def load_documents(paths: Iterable[Path]) -> list[Document]:
    documents: list[Document] = []
    for path in paths:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            documents.extend(_load_pdf(path))
        elif suffix == ".docx":
            documents.extend(_load_docx(path))
        elif suffix in TEXT_EXTENSIONS:
            documents.extend(_load_text(path))
    return documents
